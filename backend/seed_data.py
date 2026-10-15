import sys
from pathlib import Path
import openpyxl

# Set utf-8 encoding for console printing on Windows
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

from backend.database import get_db_connection, init_db

BASE_DIR = Path(__file__).resolve().parent.parent
SEED_FILE = BASE_DIR / "data" / "events_seed.xlsx"

def seed_events(force: bool = False):
    """
    Reads data/events_seed.xlsx and seeds the events table in SQLite.
    Skips if events table already contains data, unless force=True.
    Raises FileNotFoundError if data/events_seed.xlsx is missing.
    """
    init_db()

    conn = get_db_connection()
    cursor = conn.cursor()

    cursor.execute("SELECT COUNT(*) AS count FROM events;")
    row = cursor.fetchone()
    count = row["count"] if row else 0

    if count > 0 and not force:
        print(f"Events table already has {count} rows. Skipping seeding.")
        conn.close()
        return count

    if not SEED_FILE.exists():
        conn.close()
        raise FileNotFoundError(
            f"Seed file not found at '{SEED_FILE}'. Please copy the provided "
            "events_seed.xlsx file into the data/ folder before running this script."
        )

    wb = openpyxl.load_workbook(SEED_FILE)
    if "Events" not in wb.sheetnames:
        conn.close()
        raise ValueError("Worksheet 'Events' not found in events_seed.xlsx.")

    ws = wb["Events"]

    expected_headers = [
        "id", "title", "category", "description", "event_date",
        "event_time", "venue", "capacity", "fee", "icon"
    ]

    header_row = [cell.value for cell in ws[1]]
    if [h.lower() if isinstance(h, str) else h for h in header_row[:10]] != expected_headers:
        print(f"Warning: Excel header row differs from expected: {header_row}")

    rows_inserted = 0
    # Read rows from row 2 onwards
    for row_idx in range(2, ws.max_row + 1):
        row_vals = [ws.cell(row=row_idx, column=col_idx).value for col_idx in range(1, 11)]
        # Skip completely empty rows
        if not any(row_vals):
            continue

        id_val, title, category, description, event_date, event_time, venue, capacity, fee, icon = row_vals

        if not title:
            continue

        # Convert event_date to string if it was parsed as date/datetime
        if hasattr(event_date, "strftime"):
            event_date = event_date.strftime("%Y-%m-%d")
        elif event_date is not None:
            event_date = str(event_date).strip()

        # Convert event_time to string if it was parsed as time/datetime
        if hasattr(event_time, "strftime"):
            event_time = event_time.strftime("%I:%M %p")
        elif event_time is not None:
            event_time = str(event_time).strip()

        cursor.execute("""
            INSERT OR REPLACE INTO events (
                id, title, category, description, event_date, event_time, venue, capacity, fee, icon
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
        """, (
            int(id_val) if id_val is not None else None,
            str(title).strip() if title else "",
            str(category).strip() if category else "",
            str(description).strip() if description else "",
            event_date or "",
            event_time or "",
            str(venue).strip() if venue else "",
            int(capacity) if capacity is not None else 0,
            int(fee) if fee is not None else 0,
            str(icon).strip() if icon else ""
        ))
        rows_inserted += 1

    conn.commit()
    conn.close()
    print(f"Successfully seeded {rows_inserted} events into the database.")
    return rows_inserted

if __name__ == "__main__":
    seed_events()
    conn = get_db_connection()
    events = conn.execute("SELECT id, title, category, event_date, event_time, venue, capacity, fee, icon FROM events").fetchall()
    print("\n--- Current Events in Database ---")
    for ev in events:
        print(f"[{ev['id']}] {ev['icon']} {ev['title']} ({ev['category']}) | {ev['event_date']} {ev['event_time']} | {ev['venue']} | Cap: {ev['capacity']} | Fee: ₹{ev['fee']}")
    conn.close()
